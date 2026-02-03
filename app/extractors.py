from __future__ import annotations

from io import BytesIO
from typing import Callable

import pandas as pd
import pdfplumber
from docx import Document
from pptx import Presentation


TextExtractor = Callable[[bytes], str]


def _extract_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore")


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_pptx(file_bytes: bytes) -> str:
    presentation = Presentation(BytesIO(file_bytes))
    chunks: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                chunks.append(shape.text)
    return "\n".join(chunks)


def _extract_pdf(file_bytes: bytes) -> str:
    text_chunks: list[str] = []
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text_chunks.append(page.extract_text() or "")
    return "\n".join(text_chunks)


def _extract_excel(file_bytes: bytes) -> str:
    excel_data = pd.read_excel(BytesIO(file_bytes), sheet_name=None)
    chunks: list[str] = []
    for sheet_name, frame in excel_data.items():
        chunks.append(sheet_name)
        chunks.append(frame.to_csv(index=False))
    return "\n".join(chunks)


def get_extractor(filename: str) -> TextExtractor:
    extension = filename.lower().rsplit(".", maxsplit=1)[-1]
    if extension in {"txt"}:
        return _extract_txt
    if extension in {"docx", "doc"}:
        return _extract_docx
    if extension in {"pptx", "ppt"}:
        return _extract_pptx
    if extension in {"pdf"}:
        return _extract_pdf
    if extension in {"xlsx", "xls", "csv"}:
        return _extract_excel
    return _extract_txt
